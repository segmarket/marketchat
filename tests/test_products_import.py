import io
from decimal import Decimal

import pytest
from django.urls import reverse
from rest_framework_simplejwt.tokens import RefreshToken

from apps.products.models import Product
from apps.products.services.import_apply import ImportApplyError, confirm_import
from apps.products.services.import_compare import build_import_preview
from apps.products.services.import_session import create_import_session
from apps.products.services.spreadsheet import parse_csv_upload
from apps.products.services.template import build_template_workbook
from openpyxl import load_workbook
from tests.factories import ProductFactory, TenantFactory, UserFactory


def _csv_bytes(content: str) -> io.BytesIO:
    buf = io.BytesIO(content.encode("utf-8"))
    buf.name = "produtos.csv"
    return buf


def _rows_from_csv(content: str) -> list[dict]:
    return parse_csv_upload(_csv_bytes(content))


def test_template_workbook_includes_search_aliases_column():
    wb = load_workbook(io.BytesIO(build_template_workbook()), read_only=True)
    ws = wb.active
    headers = [str(cell or "").strip() for cell in next(ws.iter_rows(values_only=True))]
    wb.close()
    assert headers == ["sku", "name", "search_aliases", "price", "status"]


@pytest.mark.django_db
def test_csv_import_parses_search_aliases_column():
    rows = _rows_from_csv(
        'sku,name,search_aliases,price,status\n'
        'SKU-A,Refrigerante,"coca, cola",10.00,Ativo\n'
    )
    assert rows[0]["search_aliases"] == "coca, cola"


@pytest.mark.django_db
def test_csv_import_accepts_sinonimos_header_alias():
    rows = _rows_from_csv(
        "sku,name,sinonimos,price,status\n"
        "SKU-A,Refrigerante,coca,10.00,Ativo\n"
    )
    assert rows[0]["search_aliases"] == "coca"


@pytest.mark.django_db
def test_preview_detects_search_aliases_change():
    tenant = TenantFactory()
    ProductFactory(
        tenant=tenant,
        sku="ABC",
        name="Item ABC",
        price=Decimal("10.00"),
        status=Product.Status.ACTIVE,
        search_aliases="antigo",
    )
    rows = _rows_from_csv(
        "sku,name,search_aliases,price,status\n"
        "ABC,Item ABC,novo alias,10.00,Ativo\n"
    )
    preview = build_import_preview(tenant.id, rows)
    assert preview.updated_count == 1
    assert preview.operations[0]["search_aliases"] == "novo alias"


@pytest.mark.django_db
def test_confirm_import_persists_search_aliases(api_client):
    tenant = TenantFactory()
    user = UserFactory(tenant=tenant, email="products-aliases@example.com")
    api_client.credentials(
        HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}"
    )

    csv_content = (
        'sku,name,search_aliases,price,status\n'
        'ALIAS-1,Produto com alias,"coca, cola",9.99,Ativo\n'
    )
    preview_resp = api_client.post(
        reverse("products-upload-preview"),
        {"file": _csv_bytes(csv_content)},
        format="multipart",
    )
    token = preview_resp.json()["import_token"]
    api_client.post(
        reverse("products-upload-confirm"),
        {"import_token": token},
        format="json",
    )
    product = Product.all_objects.get(tenant=tenant, sku="ALIAS-1")
    assert product.search_aliases == "coca, cola"


@pytest.mark.django_db
def test_preview_identifies_new_skus():
    tenant = TenantFactory()
    rows = _rows_from_csv(
        "sku,name,price,status\n"
        "SKU-A,Produto A,10.00,Ativo\n"
        "SKU-B,Produto B,20.00,Ativo\n"
    )
    preview = build_import_preview(tenant.id, rows)
    assert preview.new_count == 2
    assert preview.updated_count == 0
    assert len(preview.operations) == 2
    assert all(op["action"] == "create" for op in preview.operations)


@pytest.mark.django_db
def test_preview_counts_price_changes():
    tenant = TenantFactory()
    ProductFactory(
        tenant=tenant,
        sku="ABC",
        name="Item ABC",
        price=Decimal("10.00"),
        status=Product.Status.ACTIVE,
    )
    rows = _rows_from_csv(
        "sku,name,price,status\n"
        "ABC,Item ABC,12.50,Ativo\n"
    )
    preview = build_import_preview(tenant.id, rows)
    assert preview.new_count == 0
    assert preview.updated_count == 1
    assert preview.operations[0]["action"] == "update"


@pytest.mark.django_db
def test_preview_ignores_identical_rows():
    tenant = TenantFactory()
    ProductFactory(
        tenant=tenant,
        sku="SAME",
        name="Igual",
        price=Decimal("5.00"),
        status=Product.Status.ACTIVE,
    )
    rows = _rows_from_csv(
        "sku,name,price,status\n"
        "SAME,Igual,5.00,Ativo\n"
    )
    preview = build_import_preview(tenant.id, rows)
    assert preview.new_count == 0
    assert preview.updated_count == 0
    assert preview.operations == []


@pytest.mark.django_db
def test_preview_isolates_tenants():
    tenant_a = TenantFactory()
    tenant_b = TenantFactory()
    ProductFactory(tenant=tenant_a, sku="SKU1", name="De A", price=Decimal("1.00"))

    rows = _rows_from_csv(
        "sku,name,price,status\n"
        "SKU1,De B,2.00,Ativo\n"
    )
    preview_b = build_import_preview(tenant_b.id, rows)
    assert preview_b.new_count == 1
    assert preview_b.updated_count == 0

    session = create_import_session(tenant_b.id, preview_b)
    confirm_import(session)

    product_a = Product.all_objects.get(tenant=tenant_a, sku="SKU1")
    assert product_a.name == "De A"
    assert product_a.price == Decimal("1.00")

    product_b = Product.all_objects.get(tenant=tenant_b, sku="SKU1")
    assert product_b.name == "De B"
    assert product_b.price == Decimal("2.00")


@pytest.mark.django_db
def test_upload_preview_returns_token(api_client):
    tenant = TenantFactory()
    user = UserFactory(tenant=tenant, email="products-preview@example.com")
    api_client.credentials(
        HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}"
    )

    csv_content = (
        "sku,name,price,status\n"
        "NEW-1,Novo Produto,15.00,Ativo\n"
    )
    url = reverse("products-upload-preview")
    response = api_client.post(
        url,
        {"file": _csv_bytes(csv_content)},
        format="multipart",
    )
    assert response.status_code == 200
    data = response.json()
    assert data["new_count"] == 1
    assert data["updated_count"] == 0
    assert "import_token" in data


@pytest.mark.django_db
def test_upload_confirm_applies_atomically(api_client):
    tenant = TenantFactory()
    user = UserFactory(tenant=tenant, email="products-confirm@example.com")
    api_client.credentials(
        HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}"
    )

    csv_content = (
        "sku,name,price,status\n"
        "CONF-1,Confirmado,9.99,Ativo\n"
    )
    preview_resp = api_client.post(
        reverse("products-upload-preview"),
        {"file": _csv_bytes(csv_content)},
        format="multipart",
    )
    token = preview_resp.json()["import_token"]

    confirm_resp = api_client.post(
        reverse("products-upload-confirm"),
        {"import_token": token},
        format="json",
    )
    assert confirm_resp.status_code == 200
    assert confirm_resp.json()["created"] == 1
    assert Product.all_objects.filter(tenant=tenant, sku="CONF-1").exists()

    again = api_client.post(
        reverse("products-upload-confirm"),
        {"import_token": token},
        format="json",
    )
    assert again.status_code == 409


@pytest.mark.django_db
def test_confirm_import_rejects_already_confirmed():
    tenant = TenantFactory()
    rows = _rows_from_csv("sku,name,price,status\nX,Prod,1.00,Ativo\n")
    preview = build_import_preview(tenant.id, rows)
    session = create_import_session(tenant.id, preview)
    confirm_import(session)
    with pytest.raises(ImportApplyError, match="já foi confirmada"):
        confirm_import(session)


@pytest.mark.django_db
def test_products_list_filter_by_name_sku_status(api_client):
    tenant = TenantFactory()
    user = UserFactory(tenant=tenant, email="products-filter@example.com")
    ProductFactory(tenant=tenant, sku="ARROZ-5KG", name="Arroz integral", status="active")
    ProductFactory(tenant=tenant, sku="FEIJAO-1KG", name="Feijão preto", status="inactive")

    api_client.credentials(
        HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user).access_token}"
    )
    by_sku = api_client.get(reverse("products-list"), {"sku": "arroz"})
    assert by_sku.status_code == 200
    assert len(by_sku.json()) == 1
    assert by_sku.json()[0]["sku"] == "ARROZ-5KG"

    by_name = api_client.get(reverse("products-list"), {"name": "feijão"})
    assert by_name.status_code == 200
    assert len(by_name.json()) == 1

    by_status = api_client.get(reverse("products-list"), {"status": "inactive"})
    assert by_status.status_code == 200
    assert len(by_status.json()) == 1
    assert by_status.json()[0]["status"] == "inactive"


@pytest.mark.django_db
def test_product_list_tenant_isolation(api_client):
    tenant_a = TenantFactory()
    tenant_b = TenantFactory()
    user_a = UserFactory(tenant=tenant_a, email="lista-a@example.com")
    ProductFactory(tenant=tenant_b, sku="OUTRO", name="Outro tenant")

    api_client.credentials(
        HTTP_AUTHORIZATION=f"Bearer {RefreshToken.for_user(user_a).access_token}"
    )
    response = api_client.get(reverse("products-list"))
    assert response.status_code == 200
    skus = [p["sku"] for p in response.json()]
    assert "OUTRO" not in skus
