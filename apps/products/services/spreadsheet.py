"""Leitura e normalização de planilhas de produtos."""

from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLUMNS = ("sku", "name", "price", "status")
SEARCH_ALIASES_HEADER_ALIASES = (
    "search_aliases",
    "sinonimos",
    "sinônimos",
    "sinonimo",
    "sinônimo",
)


class SpreadsheetError(ValueError):
    pass


def normalize_status(raw: Any) -> str:
    from apps.products.models import Product

    value = str(raw or "").strip().lower()
    if value in ("active", "ativo", "1", "true", "sim", "yes", "s"):
        return Product.Status.ACTIVE
    if value in ("inactive", "inativo", "0", "false", "nao", "não", "no", "n"):
        return Product.Status.INACTIVE
    if not value:
        raise SpreadsheetError("status é obrigatório.")
    raise SpreadsheetError(f"status inválido: {raw!r}")


def normalize_price(raw: Any) -> Decimal:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise SpreadsheetError("price é obrigatório.")
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)).quantize(Decimal("0.01"))
    text = str(raw).strip().replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise SpreadsheetError(f"price inválido: {raw!r}") from exc


def normalize_row(raw: dict[str, Any], *, line_number: int) -> dict[str, Any]:
    sku = str(raw.get("sku") or "").strip()
    name = str(raw.get("name") or "").strip()
    if not sku:
        raise SpreadsheetError(f"Linha {line_number}: sku é obrigatório.")
    if not name:
        raise SpreadsheetError(f"Linha {line_number}: name é obrigatório.")
    search_aliases = str(raw.get("search_aliases") or "").strip()
    return {
        "sku": sku,
        "name": name,
        "search_aliases": search_aliases,
        "price": str(normalize_price(raw.get("price"))),
        "status": normalize_status(raw.get("status")),
    }


def _canonicalize_search_aliases_header(mapping: dict[str, str]) -> None:
    if "search_aliases" in mapping:
        return
    for alias in SEARCH_ALIASES_HEADER_ALIASES:
        if alias in mapping:
            mapping["search_aliases"] = mapping[alias]
            return


def _canonicalize_search_aliases_col(col_map: dict[str, int]) -> None:
    if "search_aliases" in col_map:
        return
    for alias in SEARCH_ALIASES_HEADER_ALIASES:
        if alias in col_map:
            col_map["search_aliases"] = col_map[alias]
            return


def _normalize_headers(fieldnames: list[str] | None) -> dict[str, str]:
    if not fieldnames:
        raise SpreadsheetError("Planilha sem cabeçalho.")
    mapping: dict[str, str] = {}
    for name in fieldnames:
        key = str(name or "").strip().lower()
        if key:
            mapping[key] = name
    missing = [col for col in REQUIRED_COLUMNS if col not in mapping]
    if missing:
        raise SpreadsheetError(
            f"Colunas obrigatórias ausentes: {', '.join(missing)}. "
            f"Use: {', '.join(REQUIRED_COLUMNS)} "
            f"(opcional: search_aliases ou sinonimos)."
        )
    _canonicalize_search_aliases_header(mapping)
    return mapping


def _dedupe_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    unique: list[dict[str, Any]] = []
    for row in rows:
        sku = row["sku"]
        if sku in seen:
            continue
        seen.add(sku)
        unique.append(row)
    return unique


def parse_csv_upload(file_obj) -> list[dict[str, Any]]:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = raw.decode("utf-8-sig")
    else:
        text = str(raw)
    reader = csv.DictReader(io.StringIO(text))
    header_map = _normalize_headers(reader.fieldnames)

    rows: list[dict[str, Any]] = []
    for line_number, row in enumerate(reader, start=2):
        keyed = {col: row.get(header_map[col]) for col in REQUIRED_COLUMNS}
        if "search_aliases" in header_map:
            keyed["search_aliases"] = row.get(header_map["search_aliases"])
        if all(v is None or str(v).strip() == "" for v in keyed.values()):
            continue
        rows.append(normalize_row(keyed, line_number=line_number))

    rows = _dedupe_rows(rows)
    if not rows:
        raise SpreadsheetError("Nenhuma linha de produto encontrada na planilha.")
    return rows


def parse_xlsx_upload(file_obj) -> list[dict[str, Any]]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise SpreadsheetError("Planilha XLSX vazia.") from exc

    col_map: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = str(cell or "").strip().lower()
        if key:
            col_map[key] = idx
    missing = [col for col in REQUIRED_COLUMNS if col not in col_map]
    if missing:
        wb.close()
        raise SpreadsheetError(
            f"Colunas obrigatórias ausentes: {', '.join(missing)}."
        )
    _canonicalize_search_aliases_col(col_map)

    rows: list[dict[str, Any]] = []
    for line_number, row in enumerate(rows_iter, start=2):
        cells = list(row)
        keyed = {
            col: cells[col_map[col]] if col_map[col] < len(cells) else None
            for col in REQUIRED_COLUMNS
        }
        if "search_aliases" in col_map:
            idx = col_map["search_aliases"]
            keyed["search_aliases"] = cells[idx] if idx < len(cells) else None
        if all(v is None or str(v).strip() == "" for v in keyed.values()):
            continue
        rows.append(normalize_row(keyed, line_number=line_number))

    wb.close()
    rows = _dedupe_rows(rows)
    if not rows:
        raise SpreadsheetError("Nenhuma linha de produto encontrada na planilha.")
    return rows


def parse_upload(file_obj, *, filename: str = "") -> list[dict[str, Any]]:
    name = (filename or getattr(file_obj, "name", "") or "").lower()
    if name.endswith(".csv"):
        return parse_csv_upload(file_obj)
    if name.endswith(".xlsx"):
        return parse_xlsx_upload(file_obj)
    if name.endswith(".xls"):
        raise SpreadsheetError("Formato .xls não suportado. Use .xlsx ou .csv.")
    raise SpreadsheetError("Formato não suportado. Envie .xlsx ou .csv.")
